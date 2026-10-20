#!/usr/bin/env python3
"""
Gera o catálogo da loja (src/data/produtos.json), as imagens dos produtos
(public/img/p/*.webp) e a planilha de revisão (catalogo-revisao.xlsx)
a partir das planilhas exportadas do Mercado Livre e da Shopee.

Uso:
  python3 scripts/build_catalog.py --ml ANUNCIOS_ML.xlsx --shopee SHOPEE.xlsx --imagens "PASTA MERCADO SHOPS"

Regras:
  - Preço do site = menor preço entre Mercado Livre e Shopee (inclui o
    preço de atacado do ML para 1 unidade, quando existe).
  - Preços de atacado vêm do Mercado Livre e só entram se forem menores
    que o preço do site.
  - Peso e medidas vêm do Mercado Livre; quando não existem ou parecem
    errados, usamos um valor estimado e marcamos "estimado" na revisão.
"""
import argparse, json, os, re, zipfile, io
from PIL import Image
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# ---------------------------------------------------------------- leitura

def load_ml(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb['Anúncios']
    out = {}
    for r in list(ws.iter_rows(values_only=True))[6:]:
        if not r or not r[1]:
            continue
        item = str(r[1]).strip()
        num = lambda v: float(v) if v not in (None, '', '-') else None
        tiers = []
        for i in range(10, 20, 2):
            q, p = num(r[i]), num(r[i + 1])
            if q and p:
                tiers.append((int(q), round(p, 2)))
        out[item] = {
            'title': r[4], 'var': r[5], 'stock': num(r[6]), 'price': num(r[7]),
            'tiers': tiers, 'status': r[25],
            'dims': (num(r[26]), num(r[27]), num(r[28])), 'weight': num(r[29]),
        }
    return out


def load_shopee(path):
    # a planilha da Shopee vem com um atributo inválido que o openpyxl recusa
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename.startswith('xl/worksheets/'):
                data = re.sub(rb'\sactivePane="[^"]*"', b'', data)
            zout.writestr(it, data)
    buf.seek(0)
    ws = openpyxl.load_workbook(buf, data_only=True).active
    out = {}
    for r in ws.iter_rows(min_row=7, values_only=True):
        if not r[0] or not r[2]:
            continue
        gtin = re.sub(r'\D', '', str(r[7] or '')) or re.sub(r'\D', '', str(r[4] or ''))
        out[str(r[2]).strip()] = {
            'product': str(r[0]), 'title': r[1], 'var': r[3],
            'price': float(r[6]) if r[6] else None,
            'gtin': gtin if len(gtin) == 13 and gtin.startswith('97') else None,
            'stock': r[8],
        }
    return out

# ---------------------------------------------------------------- catálogo curado
# src: lista de ('ml', código do anúncio) e ('sh', id da variação da Shopee)
# dims: (altura, largura, comprimento) em cm; peso em kg. Só preencha para
#       corrigir ou completar o que veio do Mercado Livre.

CATS = {
    'livros': 'Livros',
    'pregacao': 'Pregação e ensino',
    'ebd': 'Revistas e apostilas',
    'hinarios': 'Hinários e corinhos',
    'agendas': 'Agendas',
    'kits': 'Kits',
}

A5_LIVRO = {'dims': (3, 17, 23), 'peso': 0.35}
A5_REVISTA = {'dims': (2, 16, 22), 'peso': 0.2}

P = [
 # ------------------------------------------------------------ LIVROS
 dict(slug='a-realidade-do-inferno', nome='A Realidade do Inferno', cat='livros', novo=True, destaque=True,
      sub='O que a Bíblia, a ciência e a experiência humana revelam sobre o lugar que ninguém quer conhecer',
      desc=['Um livro direto sobre um dos temas mais evitados da fé cristã. O Pr. Paulo Zang reúne o ensino das Escrituras, argumentos e testemunhos para responder às perguntas que muita gente tem medo de fazer: o inferno existe? Quem vai para lá? Existe saída?',
            'A mensagem final é de esperança: há salvação em Cristo para todo aquele que crê.'],
      ficha=[('Páginas', '145'), ('Formato', '14,5 x 21 cm')],
      v=[dict(src=[('ml', 'MLB4852995471'), ('sh', '209191405093')], img=['realidade do inferno.png'])]),
 dict(slug='o-pastor-de-hermas', nome='O Pastor de Hermas', cat='livros', novo=True,
      sub='Um livro do século II em nova tradução, fácil de ler',
      desc=['Escrito em Roma no século II, O Pastor de Hermas foi um dos livros mais lidos pelos primeiros cristãos, com visões, mandamentos e parábolas sobre arrependimento e vida santa.',
            'Esta edição traz uma tradução nova, feita a partir do texto grego, em linguagem fácil de ler. O texto está dividido em capítulos e versículos numerados, com notas de rodapé sobre palavras difíceis e curiosidades, introdução e glossário.'],
      v=[dict(key='simples', nome='Edição simples', src=[('ml', 'MLB7621030680'), ('sh', '179419743940')], img=['o pastor de hermas/pastor.png']),
         dict(key='luxo', nome='Edição luxo', src=[('sh', '179419743941')], img=['o pastor de hermas/pastor.png'], dims=(8, 17, 23), peso=0.8)]),
 dict(slug='a-musica-do-ceu', nome='A Música do Céu', cat='livros', novo=True,
      sub='A história de Jennie Evans Moore Seymour',
      desc=['Biografia de Jennie Evans Moore Seymour, uma das primeiras pessoas a receber o batismo no Espírito Santo no avivamento de 1906 em Los Angeles, que deu origem à Missão da Rua Azusa. Mais tarde, ela se casou com William J. Seymour, líder do avivamento.',
            'Em tom narrativo, o livro acompanha sua vida, seu ministério e seu lugar na história do Movimento Pentecostal.'],
      v=[dict(src=[('sh', '228812314782')], img=['A Musica do Céu.png'], **A5_LIVRO)]),
 dict(slug='cultos-nos-lares', nome='Cultos nos Lares', cat='livros',
      sub='Famílias que cultuam, igrejas que crescem',
      desc=['Um guia para levar o culto para dentro de casa. O livro mostra a base bíblica do culto doméstico e ajuda famílias e líderes a organizar momentos de oração, leitura da Palavra e comunhão nos lares.'],
      v=[dict(src=[('ml', 'MLB4281542619'), ('sh', '219608040821')], img=['cult.png'])]),
 dict(slug='estudos-biblicos-para-casais', nome='Estudos Bíblicos para Casais', cat='livros',
      sub='Material para líderes e encontros de casais',
      desc=['Estudos bíblicos sobre o casamento e a vida a dois, preparados para casais, pastores, professores e ministérios de casais.',
            'Pode ser usado em encontros, em classes de Escola Bíblica ou no estudo do próprio casal.'],
      ficha=[('Páginas', '196')],
      v=[dict(src=[('ml', 'MLB4281516453'), ('sh', '199607214535')], img=['Estudos Bíblicos para Casais.png'])]),
 dict(slug='orando-como-jesus-ensinou', nome='Orando como Jesus Ensinou', cat='livros',
      sub='Baseado na oração do Pai Nosso',
      desc=['Um estudo sobre a oração a partir do modelo que Jesus deixou aos discípulos. Cada parte do Pai Nosso é explicada para ajudar o leitor a orar com mais entendimento, como os discípulos pediram: "Senhor, ensina-nos a orar" (Lucas 11.1).'],
      ficha=[('Páginas', '80')],
      v=[dict(src=[('ml', 'MLB5879923530'), ('sh', '33862760240')], img=['orando.png'])]),

 # ------------------------------------------------------------ PREGAÇÃO E ENSINO
 *[dict(slug=f'mensagens-biblicas-vol-{n}', nome=f'Mensagens Bíblicas, Vol. {n}', cat='pregacao', novo=(n == 5), destaque=(n == 5),
        sub='Esboços de sermões para pregação e estudos bíblicos',
        desc=[f'Volume {n} da coleção Mensagens Bíblicas, com esboços de pregação do Pr. Paulo Zang organizados em tópicos a partir do texto bíblico.',
              'Ajuda a preparar sermões, cultos e estudos bíblicos com mais segurança e clareza.'],
        ficha=[('Encadernação', 'Espiral')],
        v=[dict(src=src, img=[img], **extra)])
   for n, src, img, extra in [
     (1, [('ml', 'MLB5881472760'), ('ml', 'MLB5879807956'), ('sh', '73851636889')], 'Mensagens biblicas.png', {'dims': (3, 18, 22), 'peso': 0.35}),
     (2, [('ml', 'MLB4422484893'), ('ml', 'MLB4422394403'), ('sh', '80957904798')], 'MENSAGENS BIBLICAS 2.png', {'dims': (3, 17, 22), 'peso': 0.35}),
     (3, [('ml', 'MLB5879923578'), ('ml', 'MLB4281182413'), ('sh', '125281092600')], 'MENBS BIBLICAS 3.png', {'dims': (3, 18, 22), 'peso': 0.35}),
     (4, [('ml', 'MLB5879872666'), ('sh', '199591688931')], 'MENSAGENS BIBLICAS 4 li.png', {}),
     (5, [('ml', 'MLB4304499185'), ('sh', '189614480784')], 'mockup1763041550446.png', {}),
   ]],
 dict(slug='estudos-biblicos-para-obreiros-vol-1', nome='Estudos Bíblicos para Obreiros, Vol. 1', cat='pregacao',
      sub='Ética, homilética e etiqueta',
      desc=['Primeiro volume da série Estudos Bíblicos para Obreiros, com estudos para a formação de obreiros e líderes da igreja, incluindo ética cristã, homilética (preparo de sermões) e etiqueta ministerial.'],
      v=[dict(src=[('ml', 'MLB5882783090'), ('ml', 'MLB5879859884'), ('sh', '73860699402')], img=['ESTUDOS BÍBLICOS PARA OBREIROS capa.png'])]),
 dict(slug='estudos-biblicos-para-obreiros-vol-2', nome='Estudos Bíblicos para Obreiros, Vol. 2', cat='pregacao',
      sub='Formação bíblica para um ministério aprovado',
      desc=['Segundo volume da série Estudos Bíblicos para Obreiros, com novos estudos para quem serve na obra de Deus: o bom obreiro, autoridade e submissão, vida devocional, dons ministeriais e o ministério pastoral.'],
      v=[dict(src=[('ml', 'MLB5879782388'), ('sh', '199587280990')], img=['ESTUDOS BÍBLICOS PARA OBREIROS  2.png'])]),
 dict(slug='estudos-biblicos-para-obreiros-vol-3', nome='Estudos Bíblicos para Obreiros, Vol. 3', cat='pregacao', novo=True,
      sub='Facilidade para manusear, clareza para aprender',
      desc=['Terceiro volume da série, em espiral, com estudos para obreiros apresentados com textos-chave, esquemas e quadros ilustrativos que facilitam o ensino e o estudo pessoal.'],
      ficha=[('Encadernação', 'Espiral')],
      v=[dict(src=[('ml', 'MLB5241017255'), ('sh', '199194392747')], img=['E B O 3/Estudos Bíblicos para Obreiros - Volume 3.png'])]),
 dict(slug='ilustracoes', nome='Ilustrações', cat='pregacao',
      sub='Histórias e anedotas cristãs para ilustrar a mensagem',
      desc=['Uma coletânea de histórias, exemplos e anedotas cristãs para enriquecer pregações, aulas de Escola Bíblica e devocionais. Ajuda o pregador a tornar a mensagem mais clara e fácil de lembrar.'],
      v=[dict(src=[('ml', 'MLB6173379276'), ('ml', 'MLB6173184186'), ('sh', '73852261816')], img=['ilustra.png'])]),
 dict(slug='voce-sabia', nome='Você Sabia?', cat='pregacao',
      sub='Mais de 1.000 curiosidades bíblicas e históricas',
      desc=['Mais de mil curiosidades sobre a Bíblia, seus personagens, lugares e costumes, e sobre a história do cristianismo.',
            'Ótimo para professores, pregadores, gincanas bíblicas e para quem gosta de aprender.'],
      v=[dict(src=[('ml', 'MLB5879730926'), ('sh', '183166117994')], img=['Curiosidades Bíblicas VOCÊ SABIA.png'])]),
 dict(slug='caderno-de-esbocos', nome='Caderno para Anotações de Esboços', cat='pregacao',
      sub='Simples e objetivo',
      desc=['Caderno para anotar e guardar seus esboços de pregação em um só lugar. Simples e objetivo, para levar ao púlpito, às reuniões e aos estudos.'],
      v=[dict(src=[('sh', '238788873685')], img=['mockup1753149006638.png'], dims=(2, 17, 23), peso=0.3)]),

 # ------------------------------------------------------------ REVISTAS E APOSTILAS
 dict(slug='o-culto-pelos-padroes-biblicos', nome='O Culto pelos Padrões Bíblicos', cat='ebd',
      sub='13 lições bíblicas',
      desc=['Estudos sobre o culto desde o Antigo Testamento até os dias atuais, mostrando como Deus deseja ser adorado.',
            'Lições: o culto pessoal e o culto coletivo, o culto no Antigo Testamento, no tabernáculo, na igreja primitiva e na história da igreja, louvor e adoração, o culto doméstico, elementos e contexto do culto, a música no culto, música cristã contemporânea, celebridades e shows gospel, o culto aceitável a Deus e o culto cristocêntrico.'],
      ficha=[('Lições', '13'), ('Páginas', '92')],
      v=[dict(src=[('ml', 'MLB4281195231'), ('sh', '63864749878')], img=['culto.png'])]),
 dict(slug='o-cristao-e-as-financas', nome='O Cristão e as Finanças', cat='ebd',
      sub='13 lições para uma vida financeira alinhada aos princípios de Deus',
      desc=['Revista para Escola Bíblica e estudos em grupo sobre dinheiro, mordomia e prosperidade à luz da Bíblia.',
            'Lições: mordomia cristã, o cristão e o dinheiro, o dízimo, o dízimo na graça, oferta aceitável, o princípio da honra, investindo em missões, mude sua mentalidade, administração, saindo das dívidas, pecados no dinheiro, finanças na família e a verdadeira prosperidade.'],
      ficha=[('Lições', '13')],
      v=[dict(src=[('sh', '24477418825')], img=['Finanças apostila.png'], **A5_REVISTA)]),
 dict(slug='o-calendario-profetico', nome='O Calendário Profético', cat='ebd',
      sub='13 lições de escatologia',
      desc=['Revista com 13 lições sobre o que a Bíblia diz a respeito do futuro e dos acontecimentos finais. Para Escola Bíblica, estudos em grupo e estudo pessoal.'],
      ficha=[('Lições', '13')],
      v=[dict(src=[('ml', 'MLB5879846756'), ('sh', '24475810036')], img=['escatologia.png'])]),
 dict(slug='o-espirito-santo', nome='O Espírito Santo', cat='ebd',
      sub='Sua pessoa, sua obra, o batismo e os dons',
      desc=['Lições bíblicas sobre a pessoa e a obra do Espírito Santo, o batismo no Espírito e os dons espirituais. Para Escola Bíblica, discipulado e estudo pessoal.'],
      ficha=[('Lições', '13')],
      v=[dict(src=[('ml', 'MLB5879923586'), ('sh', '153125013629')], img=['O ESPIRITO SANTO 2021.png'])]),
 dict(slug='desafios-da-fe-crista', nome='Desafios da Fé Cristã na Pós-Modernidade', cat='ebd',
      sub='13 lições',
      desc=['Revista com respostas bíblicas para os desafios que o cristão enfrenta hoje.',
            'Lições: pós-modernidade, tecnologia (bênção ou maldição?), depressão e ansiedade, idolatria moderna, a importância da leitura bíblica, o poder da Palavra, reacendendo a chama, relacionamentos, a família do século XXI, criação e evolução, inovação e modismos, o esfriamento do amor e os desigrejados.'],
      ficha=[('Lições', '13'), ('Formato', '14,5 x 21 cm')],
      v=[dict(src=[('sh', '192471825167')], img=['Desafios da Fé Cristã na Pos modernidade.jpg'], **A5_REVISTA)]),
 dict(slug='curso-de-batismo', nome='Curso de Batismo', cat='ebd',
      sub='Discipulado básico em 13 lições',
      desc=['Curso de discipulado para preparar novos convertidos para o batismo nas águas, em 13 lições.',
            'Existe a revista do aluno e a revista do professor. Para a turma, escolha um dos kits.'],
      ficha=[('Lições', '13')],
      v=[dict(key='aluno', nome='Revista do aluno', src=[('ml', 'MLB4371026447'), ('sh', '159789683764')], img=['curso de BATISMO/CAPA.jpg']),
         dict(key='kit-1-1', nome='Kit 1 aluno + 1 professor', src=[('sh', '238813297151')], img=['curso de BATISMO/anuncio kit.jpg', 'curso de BATISMO/Rev professor.png'], dims=(4, 16, 23), peso=0.3),
         dict(key='kit-5-1', nome='Kit 5 alunos + 1 professor', src=[('ml', 'MLB6196304508'), ('sh', '239433999912')], img=['curso de BATISMO/anuncio kit.jpg'])]),
 dict(slug='evangelismo-e-integracao', nome='Evangelismo e Integração', cat='ebd',
      sub='Estudo bíblico sobre como evangelizar e discipular',
      desc=['Estudo bíblico prático sobre evangelismo e integração de novos convertidos: como anunciar o evangelho e acompanhar quem chega à igreja até se firmar na fé.',
            'Avulso ou em kit com 10 apostilas para a equipe.'],
      v=[dict(key='unidade', nome='1 apostila', src=[('sh', '139643146385')], img=['evangelismo/1.png'], dims=(1, 16, 22), peso=0.12),
         dict(key='kit-10', nome='Kit com 10 apostilas', src=[('sh', '189590689977')], img=['evangelismo/capa anuncio.jpg'], dims=(8, 16, 22), peso=1.2)]),
 dict(slug='obreiro-aprovado', nome='O Obreiro Aprovado', cat='ebd',
      sub='Estudo bíblico para obreiros',
      desc=['Estudo bíblico para obreiros inspirado em 2 Timóteo 2.15: "Procura apresentar-te a Deus aprovado, como obreiro que não tem de que se envergonhar".',
            'Material curto e prático para reuniões de obreiros e classes de formação. Também em kit com 10 livros para a liderança.'],
      v=[dict(key='unidade', nome='1 livro', src=[('sh', '84951124969')], img=['apostila.png'], dims=(1, 16, 22), peso=0.12),
         dict(key='kit-10', nome='Kit com 10 livros', src=[('sh', '210317426288')], img=['OBREIRO APROVADO CAPA.jpg'], dims=(8, 16, 22), peso=1.2)]),
 dict(slug='olho-gordo-feiticaria-e-maldicao', nome='Olho Gordo, Feitiçaria e Maldição: Isso Pega no Crente?', cat='ebd', novo=True,
      sub='Respostas bíblicas para vencer o medo e viver seguro em Cristo',
      desc=['"Fizeram um trabalho contra você!" Poucas frases tiram tanto a paz de uma pessoa. Esta apostila responde com a Bíblia aberta se feitiçaria, olho gordo e maldição de família podem atingir o crente, começando por quem Deus é.',
            'Em quatro partes: quem é o nosso Deus, as promessas de Deus para os seus, o que são essas crenças e o que a Bíblia responde a cada uma, e como vive o crente seguro. Tem quadros para guardar, tabela comparativa e perguntas para revisão.',
            'Ideal para estudo pessoal, células, Escola Bíblica, discipulado e para líderes que orientam pessoas com medo de "trabalhos" e inveja.'],
      ficha=[('Páginas', '40'), ('Formato', 'A5 (14,8 x 21 cm)'), ('Versão bíblica', 'Almeida Revista e Corrigida')],
      v=[dict(src=[('ml', 'MLB7853245672'), ('sh', '228814739945')], img=['OLHO GORDO, INVEJA APOSTILA/TIKTOK - IMAGENS PRONTAS/01_capa_apostila_olho_gordo.jpg'])]),

 # ------------------------------------------------------------ HINÁRIOS E CORINHOS
 dict(slug='hinos-antigos', nome='Hinos Antigos', cat='hinarios', destaque=True,
      sub='200 hinos para você cantar e adorar ao Senhor',
      desc=['Livro com as letras de 200 hinos antigos que marcaram gerações de crentes, para cantar no culto, nos lares e nos círculos de oração.',
            'Escolha a versão em espiral, a versão luxo com capa dura floral ou a opção com as músicas gravadas.'],
      v=[dict(key='espiral', nome='Espiral', src=[('ml', 'MLB5879833582'), ('sh', '179725256106')], img=['mockup1771434502159.png']),
         dict(key='luxo', nome='Luxo, capa dura floral', src=[('ml', 'MLB4512962015'), ('sh', '179725256107'), ('sh', '179726136432')], img=['HINÁRIO FLORAL.png']),
         dict(key='musicas', nome='Músicas gravadas', src=[('sh', '149552371473')], img=['cartão pendrive.png', 'mockup1771434502159.png'], dims=(3, 17, 22), peso=0.35)]),
 dict(slug='hinario-sucessos-evangelicos', nome='Hinário Sucessos Evangélicos', cat='hinarios',
      sub='200 letras de canções que são sucesso até hoje',
      desc=['Hinário com as letras de 200 canções evangélicas que marcaram época e continuam sendo cantadas nas igrejas. Em espiral ou na versão luxo com capa dura floral.'],
      v=[dict(key='espiral', nome='Espiral', src=[('sh', '179407717333')], img=['sucessos evangelicos1.png'], dims=(2, 17, 22), peso=0.25),
         dict(key='luxo', nome='Luxo, capa dura floral', src=[('sh', '209187742768')], img=['mockup1775081497694.png'], dims=(3, 18, 22), peso=0.69)]),
 dict(slug='corinhos-avivados', nome='Corinhos Avivados', cat='hinarios', destaque=True,
      sub='Batei palmas, todos os povos; aclamai a Deus com voz de triunfo (Salmo 47.1)',
      desc=['Livreto com as letras dos corinhos de fogo e avivamento cantados nas igrejas pentecostais.',
            'Tamanho médio, tamanho grande com letra gigante (boa para quem tem dificuldade de leitura) ou versão luxo floral.'],
      v=[dict(key='medio', nome='Médio', src=[('ml', 'MLB4281194999'), ('sh', '199618405667')], img=['Corinhos.png']),
         dict(key='grande', nome='Grande, letra gigante', src=[('ml', 'MLB4852995791'), ('sh', '199618405666')], img=['CORINHOS LETRA GIGANTE.png']),
         dict(key='luxo', nome='Luxo floral, grande', src=[('sh', '238804206620')], img=['mockup1775274345222.png'], dims=(3, 18, 22), peso=0.69)]),

 # ------------------------------------------------------------ AGENDAS
 dict(slug='agenda-crista-feminina-2027', nome='Agenda Cristã Feminina 2027', cat='agendas', novo=True, destaque=True,
      sub='Com versículos bíblicos, wire-o e capa laminada',
      desc=['Perfeita para a organização diária e o crescimento espiritual. Tem versículos bíblicos em todas as páginas, encadernação wire-o que abre totalmente e facilita a escrita, e capa laminada de alta durabilidade.',
            'Traz calendários de 2027 e 2028. Um presente para pastoras, líderes, amigas e irmãs da igreja, ou para começar o ano com propósito e fé.'],
      ficha=[('Ano', '2027'), ('Páginas', '2 dias por página'), ('Encadernação', 'Wire-o'), ('Capa', 'Laminada')],
      extras=['2027/PROPAGANDA - ANUNCIOS/TIKTOK - IMAGENS PRONTAS/femininas/11_interna_paginas_diarias.jpg',
              '2027/PROPAGANDA - ANUNCIOS/TIKTOK - IMAGENS PRONTAS/femininas/12_interna_calendario_2027.jpg',
              '2027/PROPAGANDA - ANUNCIOS/TIKTOK - IMAGENS PRONTAS/femininas/15_interna_abertura_mes.jpg'],
      v=[dict(key=k, nome=n, src=[('ml', ml)], img=[f'2027/PROPAGANDA - ANUNCIOS/TIKTOK - IMAGENS PRONTAS/femininas/{f}.jpg'])
         for k, n, ml, f in [
           ('floral-rosa', 'Floral Rosa', 'MLB5377076841', '01_capa_floral_rosa'),
           ('floral-verde', 'Floral Verde', 'MLB5377051543', '02_capa_floral_verde'),
           ('borboletas', 'Borboletas', 'MLB7795846010', '03_capa_borboletas'),
           ('floral-lilas', 'Floral Lilás', 'MLB5377076663', '04_capa_floral_lilas'),
           ('marmore-rosa', 'Mármore Rosa', 'MLB5377076683', '05_capa_marmore_rosa'),
           ('seja-forte-e-corajosa', 'Seja Forte e Corajosa', 'MLB7795857020', '06_capa_seja_forte_e_corajosa'),
           ('sou-filha-do-rei', 'Sou Filha do Rei', 'MLB5377076839', '07_capa_sou_filha_do_rei'),
           ('leoa-forte-e-corajosa', 'Leoa Forte e Corajosa', 'MLB5377051553', '08_capa_leoa_forte_e_corajosa'),
           ('yeshua', 'Yeshua', 'MLB7795856904', '09_capa_yeshua'),
           ('mulher-virtuosa', 'Mulher Virtuosa', 'MLB7795856916', '10_capa_mulher_virtuosa'),
         ]]),
 dict(slug='agenda-crista-masculina-2027', nome='Agenda Cristã Masculina 2027', cat='agendas', novo=True, destaque=True,
      sub='Com versículos bíblicos, wire-o e capa laminada',
      desc=['Perfeita para a organização diária e o crescimento espiritual. Tem versículos bíblicos em todas as páginas, encadernação wire-o que abre totalmente e facilita a escrita, e capa laminada de alta durabilidade.',
            'Traz planejamento mensal e controle financeiro. Um presente para pastores, líderes, amigos e irmãos da igreja.'],
      ficha=[('Ano', '2027'), ('Encadernação', 'Wire-o'), ('Capa', 'Laminada')],
      extras=['2027/PROPAGANDA - ANUNCIOS/TIKTOK - IMAGENS PRONTAS/masculinas/M09_interna_paginas_diarias.jpg',
              '2027/PROPAGANDA - ANUNCIOS/TIKTOK - IMAGENS PRONTAS/masculinas/M10_interna_planejamento_mensal.jpg',
              '2027/PROPAGANDA - ANUNCIOS/TIKTOK - IMAGENS PRONTAS/masculinas/M12_interna_controle_financeiro.jpg'],
      v=[dict(key=k, nome=n, src=[('ml', ml)], img=[f'2027/PROPAGANDA - ANUNCIOS/TIKTOK - IMAGENS PRONTAS/masculinas/{f}.jpg'], dims=(8, 16, 23))
         for k, n, ml, f in [
           ('leao-seja-forte-e-corajoso', 'Leão Seja Forte e Corajoso', 'MLB7795644476', 'M01_capa_leao_seja_forte_e_corajoso'),
           ('seja-forte-e-corajoso-azul', 'Seja Forte e Corajoso Azul', 'MLB5376826061', 'M02_capa_seja_forte_e_corajoso_azul'),
           ('escudo-da-fe', 'Escudo da Fé', 'MLB7795620566', 'M03_capa_escudo_da_fe'),
           ('jesus-nome-acima-de-todo-nome', 'Jesus, Nome Acima de Todo Nome', 'MLB5376851201', 'M04_capa_jesus_nome_acima_de_todo_nome'),
           ('jesus-salva', 'Jesus Salva', 'MLB7795621194', 'M05_capa_jesus_salva'),
           ('tudo-posso', 'Tudo Posso', 'MLB5376826729', 'M06_capa_tudo_posso'),
           ('leao-aquarela', 'Leão Aquarela', 'MLB7795621150', 'M07_capa_leao_aquarela'),
           ('leao-colorido', 'Leão Colorido', 'MLB7795644472', 'M08_capa_leao_colorido'),
         ]]),
 dict(slug='agenda-crista-feminina-permanente', nome='Agenda Cristã Feminina Permanente', cat='agendas',
      sub='Sem ano impresso, com versículos e capa dura laminada',
      desc=['Agenda permanente: não tem ano impresso, então você começa a usar em qualquer data. Tem versículos bíblicos nas páginas, dois dias por página, capa dura laminada e encadernação wire-o.'],
      ficha=[('Páginas', '2 dias por página'), ('Capa', 'Dura, laminada'), ('Encadernação', 'Wire-o')],
      extras=['AGENDAS PERMANENTES/PROPAGANDA/Agenda 1 - feminina - 2DPP_PERM_1.jpg'],
      v=[dict(key=k, nome=n, src=src, img=[f'AGENDAS PERMANENTES/PROPAGANDA/mockup/{m}.png'])
         for k, n, src, m in [
           ('floral', 'Floral', [('ml', 'MLB6495334960'), ('sh', '129839125681')], 'mockup1774194140220'),
           ('floral-delicado', 'Floral Delicado', [('ml', 'MLB4549723807'), ('sh', '129839125680')], 'mockup1774193983427'),
           ('menina-forte-e-corajosa', 'Menina Forte e Corajosa', [('ml', 'MLB6495392740'), ('sh', '129839125682')], 'mockup1774193966285'),
           ('forte-e-corajosa', 'Forte e Corajosa', [('ml', 'MLB4549723809'), ('sh', '129839125683')], 'mockup1774194023143'),
           ('filha-do-rei', 'Filha do Rei', [('ml', 'MLB4549701845'), ('sh', '129839125684')], 'mockup1774193926498'),
           ('mulher-virtuosa', 'Mulher Virtuosa', [('ml', 'MLB4549723811'), ('sh', '129839125685')], 'mockup1774194070750'),
           ('orando', 'Orando', [('ml', 'MLB6495345902'), ('sh', '129839125686')], 'mockup1774194103162'),
         ]]),
 dict(slug='agenda-crista-permanente-diaria', nome='Agenda Cristã Permanente Diária', cat='agendas',
      sub='Um dia por página, com versículo bíblico em cada dia',
      desc=['Para quem escreve bastante: uma página inteira para cada dia, com um versículo bíblico no rodapé. É permanente, sem ano impresso, para começar a usar em qualquer data.'],
      ficha=[('Páginas', '1 dia por página'), ('Encadernação', 'Wire-o')],
      extras=['AGENDAS PERMANENTES/PROPAGANDA/Floral Rosa 1DPP_permanente.jpg', 'AGENDAS PERMANENTES/PROPAGANDA/Floral Rosa 1DPP_permanente2.jpg'],
      v=[dict(key=k, nome=n, src=[('ml', ml)], img=[f'AGENDAS PERMANENTES/PROPAGANDA/mockup/{m}.png'])
         for k, n, ml, m in [
           ('forte-e-corajosa', 'Forte e Corajosa', 'MLB7439411022', 'mockup1774194023143'),
           ('forte-e-corajosa-floral', 'Forte e Corajosa Floral', 'MLB7439411024', 'mockup1774193966285'),
           ('floral-delicado', 'Floral Delicado', 'MLB7439366436', 'mockup1774193983427'),
           ('orando', 'Orando', 'MLB5081041495', 'mockup1774194103162'),
         ]]),
 dict(slug='agenda-crista-masculina-permanente', nome='Agenda Cristã Masculina Permanente', cat='agendas',
      sub='Sem ano impresso, com versículos e capa laminada',
      desc=['Agenda permanente masculina com versículos bíblicos. Não tem ano impresso, então você começa a usar em qualquer data.'],
      v=[dict(src=[('sh', '159792684591')], img=['AGENDAS PERMANENTES/CAPAS/Forte e Corajoso.jpg'], dims=(3, 17, 22), peso=0.69)]),

 # ------------------------------------------------------------ KITS
 dict(slug='kit-estudos-biblicos-para-obreiros', nome='Kit Estudos Bíblicos para Obreiros, Vols. 1 e 2', cat='kits', destaque=True,
      sub='Material de apoio para o ministério',
      desc=['Os dois primeiros volumes da série Estudos Bíblicos para Obreiros juntos.',
            'Temas: o bom obreiro, cuidado de si e da doutrina, autoridade e submissão, vida devocional e ministério, sacerdócio e dons ministeriais, ministério pastoral, deveres e tipos de obreiros, diáconos, diaconisas e presbíteros, e o ministério do pregador.'],
      v=[dict(src=[('ml', 'MLB6630929094'), ('sh', '199188218412')], img=['image_w1024_h1024_capa-shopee-kit-obreiros-var1.jpeg'])]),
 dict(slug='kit-mensagens-biblicas', nome='Kit Mensagens Bíblicas', cat='kits',
      sub='Coleção de esboços de pregação do Pr. Paulo Zang',
      desc=['Monte sua coleção de esboços de pregação com os volumes de Mensagens Bíblicas em kit, por um preço menor do que comprando separado.'],
      v=[dict(key='vol-1-a-5', nome='Vols. 1 a 5', src=[('ml', 'MLB4322212075'), ('sh', '199614919581')], img=['kit mens bib 1 a 5.jpg']),
         dict(key='vol-1-a-4', nome='Vols. 1 a 4', src=[('ml', 'MLB5879859938'), ('sh', '209593250469')], img=['kit mensagens bíblicas.jpg']),
         dict(key='vol-1-a-3', nome='Vols. 1 a 3', src=[('sh', '125396832710')], img=['0000 KIT 3 LIVROS ML.jpg'], dims=(6, 17, 23), peso=0.9),
         dict(key='vol-4-e-5', nome='Vols. 4 e 5', src=[('sh', '239439617990')], img=['@compor:MENSAGENS BIBLICAS 4 li.png|mockup1763041550446.png'], dims=(4, 17, 23), peso=0.7)]),
 dict(slug='kit-hinos-sucessos-corinhos', nome='Kit Hinos Antigos, Sucessos Evangélicos e Corinhos Avivados', cat='kits',
      sub='Os três livros de louvor da editora',
      desc=['Hinos Antigos, Hinário Sucessos Evangélicos e Corinhos Avivados juntos. Escolha a versão simples ou a versão luxo com capas florais.'],
      v=[dict(key='simples', nome='Versão simples', src=[('ml', 'MLB4283474631'), ('sh', '239427199797')], img=['@compor:mockup1771434502159.png|sucessos evangelicos1.png|Corinhos.png']),
         dict(key='luxo', nome='Versão luxo floral', src=[('ml', 'MLB4612441885'), ('sh', '199187830420')], img=['kIT fLORAL.jpg'])]),
 dict(slug='kit-corinhos-avivados', nome='Kit Corinhos Avivados', cat='kits',
      sub='Para distribuir na igreja, no círculo de oração e no louvor',
      desc=['Corinhos Avivados em quantidade, para a igreja, o círculo de oração ou o grupo de louvor.'],
      v=[dict(key='3', nome='3 unidades', src=[('sh', '199598862099')], img=['@compor:Corinhos.png|Corinhos.png|Corinhos.png'], dims=(5, 11, 20), peso=0.9),
         dict(key='10', nome='10 unidades', src=[('ml', 'MLB4281202827'), ('sh', '199163434757')], img=['10 UNIDADES CORINHOS.jpg']),
         dict(key='20', nome='20 unidades', src=[('sh', '164915846856')], img=['CORINHOS ATACADO.jpg'], dims=(12, 17, 22), peso=4.2)]),
 dict(slug='box-pregador', nome='Box Pregador', cat='kits',
      sub='5 livros e, de brinde, o curso básico O Pregador e a Pregação',
      desc=['Kit com cinco livros para quem prega e ensina a Palavra, acompanhado do curso básico O Pregador e a Pregação como brinde.'],
      v=[dict(src=[('sh', '116377164937')], img=['KIT PREGADOR.jpg'], dims=(12, 17, 23), peso=1.6)]),
]

# ---------------------------------------------------------------- imagens

def open_rgba(path):
    im = Image.open(path)
    im.load()
    return im.convert('RGBA')


def trim(im):
    a = im.split()[3]
    if a.getextrema()[0] < 250:  # tem transparência
        box = a.point(lambda v: 255 if v > 40 else 0).getbbox()
        return im.crop(box), 'transparente'
    rgb = im.convert('RGB')
    w, h = rgb.size
    corners = [rgb.getpixel((2, 2)), rgb.getpixel((w - 3, 2)), rgb.getpixel((2, h - 3)), rgb.getpixel((w - 3, h - 3))]
    white = all(min(c) > 235 for c in corners)
    if white:
        gray = rgb.convert('L').point(lambda v: 255 if v < 238 else 0)
        box = gray.getbbox()
        if box:
            pad = int(max(w, h) * 0.02)
            box = (max(0, box[0] - pad), max(0, box[1] - pad), min(w, box[2] + pad), min(h, box[3] + pad))
            return im.crop(box), 'branco'
    return im, 'foto'


def compose(paths, src_root):
    ims = [trim(open_rgba(os.path.join(src_root, p)))[0] for p in paths]
    H = 900
    ims = [i.resize((int(i.width * H / i.height), H), Image.LANCZOS) for i in ims]
    step = int(ims[0].width * 0.55)
    W = step * (len(ims) - 1) + ims[-1].width
    canvas = Image.new('RGBA', (W, H + 40), (0, 0, 0, 0))
    for n, i in enumerate(ims):
        canvas.alpha_composite(i, (n * step, 40 if n % 2 == 0 else 0))
    return canvas


def save_image(spec, src_root, out_dir, name, cache):
    if spec in cache:
        return cache[spec]
    if spec.startswith('@compor:'):
        im, kind = compose(spec[8:].split('|'), src_root), 'transparente'
    else:
        im, kind = trim(open_rgba(os.path.join(src_root, spec)))
    files = {}
    for size in (480, 960):
        c = im.copy()
        c.thumbnail((size, size), Image.LANCZOS)
        fn = f'{name}-{size}.webp'
        if kind == 'transparente':
            c.save(os.path.join(out_dir, fn), 'WEBP', quality=82, method=6)
        else:
            c.convert('RGB').save(os.path.join(out_dir, fn), 'WEBP', quality=80, method=6)
        files[size] = (fn, c.size)
    res = {'src': f'/img/p/{files[960][0]}', 'thumb': f'/img/p/{files[480][0]}',
           'w': files[960][1][0], 'h': files[960][1][1], 'fundo': kind}
    cache[spec] = res
    return res

# ---------------------------------------------------------------- montagem

def build(ml, sh, src_root):
    out_dir = os.path.join(ROOT, 'public', 'img', 'p')
    os.makedirs(out_dir, exist_ok=True)
    for f in os.listdir(out_dir):
        os.remove(os.path.join(out_dir, f))
    cache, produtos, revisao, faltas = {}, [], [], []
    for p in P:
        variantes = []
        for i, v in enumerate(p['v']):
            key = v.get('key')
            vid = p['slug'] + (f'--{key}' if key else '')
            precos, tiers, fontes, ml_rows = [], [], [], []
            gtin = None
            for kind, code in v['src']:
                if kind == 'ml':
                    row = ml.get(code)
                    if not row:
                        faltas.append(f'{vid}: anúncio {code} não está na planilha do Mercado Livre'); continue
                    ml_rows.append(row)
                    if row['price']:
                        precos.append(row['price']); fontes.append(('ML', row['price']))
                    for q, pr in row['tiers']:
                        if q <= 1:
                            precos.append(pr); fontes.append(('ML atacado 1 un.', pr))
                    if not tiers:
                        tiers = [t for t in row['tiers'] if t[0] > 1]
                else:
                    row = sh.get(code)
                    if not row:
                        faltas.append(f'{vid}: variação {code} não está na planilha da Shopee'); continue
                    if row['price']:
                        precos.append(row['price']); fontes.append(('Shopee', row['price']))
                    gtin = gtin or row['gtin']
            if not precos:
                faltas.append(f'{vid}: sem preço nas planilhas'); continue
            preco = round(min(precos), 2)
            atacado, seen = [], set()
            for q, pr in sorted(tiers):
                if pr < preco and q not in seen:
                    atacado.append({'min': q, 'preco': pr}); seen.add(q)
            estimado = 'dims' in v or 'peso' in v
            if 'dims' in v and 'peso' not in v and ml_rows:
                peso = ml_rows[0]['weight']
                estimado = False if peso else True
            else:
                peso = v.get('peso') or (ml_rows[0]['weight'] if ml_rows else None)
            dims = v.get('dims') or (ml_rows[0]['dims'] if ml_rows else None)
            if not dims or not peso:
                faltas.append(f'{vid}: sem peso/medidas'); continue
            a, l, c = sorted([float(x) for x in dims])  # menor = altura (espessura)
            imgs = [save_image(s, src_root, out_dir, f'{p["slug"]}{"-" + key if key else ""}-{n}', cache) for n, s in enumerate(v['img'])]
            variantes.append({
                'id': vid, 'nome': v.get('nome'), 'preco': preco, 'atacado': atacado,
                'peso': round(float(peso), 3), 'dims': {'altura': a, 'largura': l, 'comprimento': c},
                'gtin': gtin, 'imagens': imgs,
            })
            revisao.append([p['nome'], v.get('nome') or '', CATS[p['cat']], preco,
                            ', '.join(f'{f}: {pr:.2f}'.replace('.', ',') for f, pr in fontes),
                            '; '.join(f'{t["min"]}+ un.: {t["preco"]:.2f}'.replace('.', ',') for t in atacado),
                            float(peso), f'{c:g} x {l:g} x {a:g}', 'sim' if estimado else '', gtin or ''])
        if not variantes:
            continue
        extras = [save_image(s, src_root, out_dir, f'{p["slug"]}-extra-{n}', cache) for n, s in enumerate(p.get('extras', []))]
        produtos.append({
            'slug': p['slug'], 'nome': p['nome'], 'categoria': p['cat'], 'subtitulo': p.get('sub'),
            'descricao': p['desc'], 'ficha': [list(x) for x in p.get('ficha', [])],
            'novo': bool(p.get('novo')), 'destaque': bool(p.get('destaque')),
            'variantes': variantes, 'extras': extras,
            'precoMin': min(x['preco'] for x in variantes), 'precoMax': max(x['preco'] for x in variantes),
        })
    return produtos, revisao, faltas


def write_review(rows, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Catálogo'
    head = ['Produto', 'Variação', 'Categoria', 'Preço no site (R$)', 'Preços encontrados', 'Atacado (ML)',
            'Peso (kg)', 'Medidas C x L x A (cm)', 'Peso/medida estimados?', 'ISBN']
    ws.append(head)
    for r in rows:
        ws.append(r)
    from openpyxl.styles import Font, PatternFill, Alignment
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='14213D')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    widths = [44, 30, 20, 14, 40, 34, 10, 20, 14, 16]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = '#,##0.00'
        row[6].number_format = '0.00'
        if row[8].value == 'sim':
            for c in row:
                c.fill = PatternFill('solid', fgColor='FFF4D6')
    ws.freeze_panes = 'A2'
    wb.save(path)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--ml', required=True)
    ap.add_argument('--shopee', required=True)
    ap.add_argument('--imagens', required=True)
    a = ap.parse_args()
    ml, sh = load_ml(a.ml), load_shopee(a.shopee)
    produtos, revisao, faltas = build(ml, sh, a.imagens)
    data = {'categorias': [{'id': k, 'nome': v} for k, v in CATS.items()], 'produtos': produtos}
    os.makedirs(os.path.join(ROOT, 'src', 'data'), exist_ok=True)
    with open(os.path.join(ROOT, 'src', 'data', 'produtos.json'), 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    write_review(revisao, os.path.join(ROOT, 'catalogo-revisao.xlsx'))
    print(f'{len(produtos)} produtos, {sum(len(p["variantes"]) for p in produtos)} variações')
    for x in faltas:
        print('ATENÇÃO:', x)
